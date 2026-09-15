"""Formatting that hides text in Word and Excel files.

Neither format has a fixed rendering to check pixels against, so visibility is
read from the markup. Each reason is a fact about the file and is reported as
confirmed.

- Word: text formatted as hidden (directly or through its style), white text
  with no shading or page colour behind it, and text below the minimum size.
- Excel: hidden and very hidden sheets, hidden rows and columns, cells whose
  number format displays nothing, and white text on an unfilled cell.
"""

from __future__ import annotations

import contextlib
import zipfile
from pathlib import Path
from typing import Any

from complydoc.config.schema import VisibilityConfig
from complydoc.hidden.visibility import HiddenRun

__all__ = ["docx_hidden_runs", "xlsx_hidden_runs"]

_MAX_TEXT = 2000


def _letters(text: str) -> int:
    return sum(c.isalnum() for c in text)


def _is_white(value: str | None, config: VisibilityConfig) -> bool:
    if not value or len(value) < 6:
        return False
    try:
        channels = [int(value[-6:][i : i + 2], 16) for i in (0, 2, 4)]
    except ValueError:
        return False
    return min(channels) >= int(config.near_white * 255)


# --------------------------------------------------------------------------- Word


def docx_hidden_runs(
    path: Path, config: VisibilityConfig
) -> tuple[list[HiddenRun], list[str], bool]:
    import docx
    from docx.opc.exceptions import PackageNotFoundError
    from docx.oxml.ns import qn

    try:
        source = docx.Document(str(path))
    except (PackageNotFoundError, zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
        return [], [f"the file could not be opened to check for hidden text ({exc})"], False

    styles = _styles(source, qn, config)
    background = source.element.find(qn("w:background"))
    coloured_page = background is not None and not _is_white(background.get(qn("w:color")), config)

    runs: list[HiddenRun] = []
    for paragraph in source.element.body.iter(qn("w:p")):
        paragraph_props = paragraph.find(qn("w:pPr"))
        paragraph_style = _style_id(paragraph_props, "w:pStyle", qn)
        shaded_paragraph = _shaded(paragraph_props, qn, config)

        current: tuple[str, ...] = ()
        parts: list[str] = []

        for run in paragraph.iter(qn("w:r")):
            text = "".join(node.text or "" for node in run.iter(qn("w:t")))
            if not text:
                continue
            run_props = run.find(qn("w:rPr"))
            effective: dict[str, Any] = {}
            for style_id in (paragraph_style, _style_id(run_props, "w:rStyle", qn)):
                if style_id:
                    effective.update(styles.get(style_id, {}))
            if run_props is not None:
                effective.update(_properties(run_props, qn, config))

            reasons: list[str] = []
            if effective.get("vanish"):
                reasons.append("formatted as hidden text")
            behind = shaded_paragraph or effective.get("shaded") or coloured_page
            if _is_white(effective.get("color"), config) and not behind:
                reasons.append("white text")
            size = effective.get("size")
            if size is not None and size < config.min_font_size_pt:
                reasons.append(f"font size below {config.min_font_size_pt:g}pt")

            if tuple(reasons) != current:
                _flush(runs, current, parts, config)
                current, parts = tuple(reasons), []
            parts.append(text)
        _flush(runs, current, parts, config)

    return runs, [], True


def _flush(
    runs: list[HiddenRun], reasons: tuple[str, ...], parts: list[str], config: VisibilityConfig
) -> None:
    text = " ".join("".join(parts).split())
    if reasons and _letters(text) >= config.min_characters:
        runs.append(HiddenRun(None, text[:_MAX_TEXT], "confirmed", list(reasons)))


def _style_id(props: Any, tag: str, qn: Any) -> str | None:
    if props is None:
        return None
    element = props.find(qn(tag))
    return element.get(qn("w:val")) if element is not None else None


def _shaded(props: Any, qn: Any, config: VisibilityConfig) -> bool:
    if props is None:
        return False
    shading = props.find(qn("w:shd"))
    if shading is None:
        return False
    fill = shading.get(qn("w:fill"))
    return fill not in (None, "auto") and not _is_white(fill, config)


def _properties(props: Any, qn: Any, config: VisibilityConfig) -> dict[str, Any]:
    found: dict[str, Any] = {}
    vanish = props.find(qn("w:vanish"))
    if vanish is not None:
        found["vanish"] = vanish.get(qn("w:val")) not in ("0", "false", "off")
    color = props.find(qn("w:color"))
    if color is not None:
        found["color"] = color.get(qn("w:val"))
    size = props.find(qn("w:sz"))
    if size is not None:
        with contextlib.suppress(TypeError, ValueError):
            found["size"] = int(size.get(qn("w:val"))) / 2
    if props.find(qn("w:highlight")) is not None or _shaded(props, qn, config):
        found["shaded"] = True
    return found


def _styles(source: Any, qn: Any, config: VisibilityConfig) -> dict[str, dict[str, Any]]:
    """Run properties per style id, with each style's base styles applied first."""
    own: dict[str, dict[str, Any]] = {}
    based_on: dict[str, str | None] = {}
    try:
        elements = source.styles.element.findall(qn("w:style"))
    except (AttributeError, KeyError, ValueError):
        return {}
    for style in elements:
        style_id = style.get(qn("w:styleId"))
        if not style_id:
            continue
        props = style.find(qn("w:rPr"))
        own[style_id] = _properties(props, qn, config) if props is not None else {}
        base = style.find(qn("w:basedOn"))
        based_on[style_id] = base.get(qn("w:val")) if base is not None else None

    resolved: dict[str, dict[str, Any]] = {}

    def resolve(style_id: str, depth: int = 0) -> dict[str, Any]:
        if style_id in resolved:
            return resolved[style_id]
        parent = based_on.get(style_id)
        merged = dict(resolve(parent, depth + 1)) if parent in own and depth < 20 else {}
        merged.update(own.get(style_id, {}))
        resolved[style_id] = merged
        return merged

    for style_id in own:
        resolve(style_id)
    return resolved


# -------------------------------------------------------------------------- Excel


def xlsx_hidden_runs(
    path: Path, config: VisibilityConfig
) -> tuple[list[HiddenRun], list[str], bool]:
    import openpyxl
    from openpyxl.utils import column_index_from_string
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = openpyxl.load_workbook(str(path), data_only=True, keep_links=False)
    except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
        return [], [f"the file could not be opened to check for hidden text ({exc})"], False

    runs: list[HiddenRun] = []
    try:
        for index, sheet in enumerate(workbook.worksheets):
            number = index + 1
            title = sheet.title
            state = getattr(sheet, "sheet_state", "visible")
            if state in ("hidden", "veryHidden"):
                values = [
                    str(cell.value)
                    for row in sheet.iter_rows()
                    for cell in row
                    if cell.value is not None and str(cell.value).strip()
                ]
                kind = "very hidden" if state == "veryHidden" else "hidden"
                _add(runs, number, values, f"{kind} sheet {title!r}", config)
                continue

            hidden_rows = {i for i, dimension in sheet.row_dimensions.items() if dimension.hidden}
            hidden_columns: set[int] = set()
            for key, dimension in sheet.column_dimensions.items():
                if dimension.hidden:
                    low = dimension.min or column_index_from_string(key)
                    high = dimension.max or low
                    hidden_columns.update(range(low, high + 1))

            groups: dict[str, list[str]] = {}
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None or not str(cell.value).strip():
                        continue
                    reason = None
                    if cell.row in hidden_rows:
                        reason = f"hidden row in sheet {title!r}"
                    elif cell.column in hidden_columns:
                        reason = f"hidden column in sheet {title!r}"
                    elif _format_shows_nothing(cell.number_format):
                        reason = f"number format that displays nothing, in sheet {title!r}"
                    elif _white_on_nothing(cell, config):
                        reason = f"white text in sheet {title!r}"
                    if reason:
                        groups.setdefault(reason, []).append(str(cell.value))
            for reason, values in groups.items():
                _add(runs, number, values, reason, config)
    finally:
        with contextlib.suppress(Exception):
            workbook.close()
    return runs, [], True


def _add(
    runs: list[HiddenRun], page: int, values: list[str], reason: str, config: VisibilityConfig
) -> None:
    text = " | ".join(" ".join(value.split()) for value in values)
    if _letters(text) >= config.min_characters:
        runs.append(HiddenRun(page, text[:_MAX_TEXT], "confirmed", [reason]))


def _format_shows_nothing(number_format: str | None) -> bool:
    shown = str(number_format or "")
    return ";" in shown and not shown.replace(";", "").strip()


def _white_on_nothing(cell: Any, config: VisibilityConfig) -> bool:
    color = getattr(cell.font, "color", None)
    rgb = getattr(color, "rgb", None) if color is not None else None
    if not isinstance(rgb, str) or not _is_white(rgb, config):
        return False
    fill = cell.fill
    if getattr(fill, "fill_type", None) in (None, "none"):
        return True
    background = getattr(getattr(fill, "fgColor", None), "rgb", None)
    return isinstance(background, str) and _is_white(background, config)
